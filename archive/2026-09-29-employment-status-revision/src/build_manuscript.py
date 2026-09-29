"""Build the manuscript from one text source and computed results.

No network calls and no manual numerical transcription in the principal result.
Word pagination is renderer-dependent; the distributed file is separately
rendered and checked. The editable source uses [@key] citation markers.
"""
from pathlib import Path
import json,re,datetime,zipfile,io
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]

def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'244D5B');pr.append(col)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'21');pr.append(size)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for cell,label in zip(t.rows[0].cells,headers):cell.text=label
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CFD5D9');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for i,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if i==0:rp.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=(i<len(t.rows)-1)
                for run in p.runs:run.font.size=Pt(10.5);run.bold=(i==0)
    return t

def main():
    r=json.loads((ROOT/'results/results.json').read_text())
    refs={x['key']:x for x in json.loads((ROOT/'manuscript/references.json').read_text(encoding='utf8'))}
    values={'N':f"{r['models']['adjusted']['n']:,}"}
    for label,model in [('A','adjusted'),('U','unadjusted'),('D','no_disability')]:
        for suffix,k in [('PR','ratio'),('LO','lower'),('HI','upper')]:values[label+suffix]=f"{r['models'][model][k]:.2f}"
    for g in [0,1]:
        e=next(x for x in r['prevalence'] if x['year']=='Pooled' and x['outcome']=='outcome' and x['exposed']==g)
        for suffix,k in [('', 'prevalence'),('LO','lower'),('HI','upper')]:values[f'P{g}'+suffix]=f"{100*e[k]:.1f}"
    source=(ROOT/'manuscript/paper.md').read_text(encoding='utf8')
    for key,value in values.items():source=source.replace('{{'+key+'}}',value)
    assert not re.search(r'\{\{.+?\}\}',source),'Unresolved numerical token'
    order=[]
    def cite(match):
        keys=[x.strip().lstrip('@') for x in match.group(1).split(';')]
        for key in keys:
            assert key in refs,key
            if key not in order:order.append(key)
        return '['+', '.join(str(order.index(k)+1) for k in keys)+']'
    source=re.sub(r'\[(@[^\]]+)\]',cite,source)
    assert len(order)==len(refs),'Uncited reference'
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.85);sec.bottom_margin=Inches(.85);sec.left_margin=Inches(1);sec.right_margin=Inches(1)
    sec.footer_distance=Inches(.42)
    # Clear decorative theme residue from the installed default template.
    for element in list(doc.styles.element.iter()):
        if element.tag in [qn('w:pBdr'),qn('w:spacing')] and element.getparent().tag==qn('w:rPr'):
            element.getparent().remove(element)
        elif element.tag==qn('w:pBdr'):
            element.getparent().remove(element)
        elif element.tag==qn('w:rFonts'):
            for attr in list(element.attrib):
                if attr.endswith('Theme'):del element.attrib[attr]
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5)
    normal.font.color.rgb=RGBColor.from_string('111111')
    normal.paragraph_format.line_spacing=1.14;normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',23),('Subtitle',13),('Heading 1',14),('Heading 2',12),('Caption',11)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('111111')
        s.font.bold=name not in ['Subtitle'];s.font.italic=False
        s.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run();run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    doc.core_properties.title='Job seeking as a psychiatric exposure'
    doc.core_properties.subject='A public health position from the hiring industry'
    doc.core_properties.author='Jacob E. Thomas'
    doc.core_properties.keywords='job seeking; unemployment; mental health; structural exposure'
    doc.core_properties.created=datetime.datetime(2026,9,29,0,0,0)
    doc.core_properties.modified=datetime.datetime(2026,9,29,0,0,0)
    md=[];front=True
    def para(text,style=None):
        rendered=re.sub(r'(?<=\d)–(?=\d)','\u2011',text)
        p=doc.add_paragraph(rendered,style=style);md.append(text);return p
    def note(text):
        p=para(text)
        p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(9)
        for run in p.runs:run.font.size=Pt(10)
    def mtable(headers,rows):
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join('---' for _ in headers)+' |',*['| '+' | '.join(str(x) for x in row)+' |' for row in rows]])
    for block in source.strip().split('\n\n'):
        if block=='@@PAGEBREAK@@':doc.add_page_break();front=False;continue
        if block=='@@FIGURE1@@':
            para('Figure 1. Symptom burden by employment category','Caption')
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(ROOT/'results/figure1.png'),width=Inches(6.5))
            shape._inline.docPr.set('descr',f"Weighted prevalence was {values['P1']} percent in the unemployment/search category and {values['P0']} percent in the employed category. Whiskers show 95 percent confidence intervals.")
            md.append('![Figure 1](../results/figure1.png)')
            note('Points show pooled survey-weighted prevalence; whiskers show 95% confidence intervals. Adults aged 18–64 in the 2022 and 2025 NHIS, complete-case sample (n = 26,281). The unemployment/search category is not a verified measure of active searching. Symptoms are PHQ-8 or GAD-7 scores of at least 10, not clinical diagnoses.')
            continue
        if block=='@@TABLE1@@':
            para('Table 1. Prevalence ratios and sensitivity analyses','Caption')
            rows=[]
            for label,key in [('Primary outcome, unadjusted','unadjusted'),('Primary outcome, adjusted','adjusted'),('Depressive symptoms, adjusted','depression'),('Anxiety symptoms, adjusted','anxiety'),('All symptom items complete, adjusted','all_items'),('Narrower employed comparator, adjusted','narrow_comparator'),('Adults without disability, adjusted','no_disability')]:
                m=r['models'][key];rows.append([label,f"{m['n']:,}",f"{m['ratio']:.2f} ({m['lower']:.2f}–{m['upper']:.2f})"])
            headers=['Analysis','Sample n','Prevalence ratio (95% CI)']
            table(doc,headers,rows,[3.4,.8,2.3]);mtable(headers,rows)
            note('Each ratio compares the unemployment/search group with the employed group. Primary outcome: moderate or severe anxiety or depressive symptoms. Adjusted models include year, age group, sex, race and Hispanic origin, education, and Census region. The narrower comparator requires work last week or temporary absence from a job. Disability exclusion uses DISAB3_A = 2. Sample sizes are unweighted; all estimates account for survey weights, strata, and clustering.')
            continue
        if block=='@@TABLE2@@':
            para('Table 2. Hiring practices that can be measured and changed','Caption')
            headers=['Priority','Observable measure','Proposed change to evaluate']
            rows=[['Vacancy accuracy','Whether an advertised vacancy is current, authorized, and accepting applications; time to remove closed listings.','Label vacancy status, distinguish ongoing talent pools, and close unavailable roles promptly.'],['Predictable communication','Days without an update; whether stages and decision dates are stated; whether applicants receive closure.','Publish the process and update schedule; provide accurate status and a defined closure procedure.'],['Proportionate demands','Applicant time and out-of-pocket costs; repeated forms; length and number of unpaid assessments.','Remove duplicate demands, limit unpaid tasks, and test compensation for substantial assessments.'],['Accessible review','Availability and use of accommodations, explanations of process, and review of disputed exclusions.','Offer accessible routes to request assistance or review; audit differences in access and outcomes.']]
            table(doc,headers,rows,[1.1,2.65,2.75]);mtable(headers,rows)
            note('These are proposed intervention targets, not demonstrated causes of psychiatric disorder or proven treatments. Evaluate mental health, employment quality, access, and unintended effects together. Process transparency does not require sharing confidential evaluations or collecting clinical information for hiring decisions.')
            continue
        if block=='@@REFERENCES@@':
            doc.add_page_break();para('References','Heading 1')
            for n,key in enumerate(order,1):
                ref=refs[key];p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.23);p.paragraph_format.first_line_indent=Inches(-.23)
                p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
                p.add_run(f"{n}. {ref['text']} ");hyperlink(p,ref['url'],ref['url'])
                for run in p.runs:run.font.size=Pt(10.5)
                md.append(f"{n}. {ref['text']} {ref['url']}")
            continue
        if block.startswith('# '):para(block[2:],'Title');continue
        if block.startswith('## '):para(block[3:],'Subtitle' if front else 'Heading 1');continue
        if block.startswith('### '):para(block[4:],'Heading 2');continue
        p=para(block)
        if front and not block.startswith('The health consequences'):
            p.paragraph_format.space_after=Pt(5)
            if block.startswith(('Correspondence:','Article type:','Keywords:')):
                for run in p.runs:run.font.size=Pt(10)
    target=ROOT/'manuscript/Job_seeking_as_a_psychiatric_exposure.docx'
    memory=io.BytesIO();doc.save(memory)
    # Remove ZIP timestamp variability. All substantive content is deterministic.
    with zipfile.ZipFile(memory) as original,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as out:
        for name in original.namelist():
            info=zipfile.ZipInfo(name,date_time=(2026,9,29,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
            out.writestr(info,original.read(name))
    (ROOT/'manuscript/manuscript.md').write_text('\n\n'.join(md)+'\n',encoding='utf8',newline='\n')
    (ROOT/'audit/citation_order.json').write_text(json.dumps(order,indent=2)+'\n',newline='\n')
    print(f'Created {target.name}; {len(order)} references; approximately {len(" ".join(md).split())} words including references and tables')

if __name__=='__main__':main()
