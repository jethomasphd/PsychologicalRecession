"""Build the opinion editorial from one text source and existing computed results."""
from pathlib import Path
import datetime,io,json,re,zipfile
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]

def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'244D5B');pr.append(col)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'21');pr.append(size)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def main():
    result=json.loads((ROOT/'results/results.json').read_text(encoding='utf8'))
    models=pd.read_csv(ROOT/'results/models.csv');annual=pd.read_csv(ROOT/'results/annual_burden.csv')
    refs={r['key']:r for r in json.loads((ROOT/'manuscript/references.json').read_text(encoding='utf8'))}
    values={'N':f"{result['primary_n']:,}",'SY':str(result['primary_state_years'])}
    for token,label,term in [('H','Primary','lower_hiring'),('C','Competition per opening','competition')]:
        row=models[(models.model==label)&(models.term==term)].iloc[0]
        for suffix,field in [('', 'estimate'),('LO','lower'),('HI','upper')]:values[token+suffix]=f'{row[field]:.2f}'
    for year in [2013,2024]:values['P'+str(year)]=f"{annual[(annual.year==year)&(annual.group=='Out of work')].iloc[0].fmd_percent:.1f}"
    source=(ROOT/'manuscript/paper.md').read_text(encoding='utf8')
    for key,value in values.items():source=source.replace('{{'+key+'}}',value)
    assert not re.search(r'\{\{.+?\}\}',source)
    order=[]
    def cite(match):
        keys=[k.strip().lstrip('@') for k in match.group(1).split(';')]
        for key in keys:
            assert key in refs,key
            if key not in order:order.append(key)
        return '['+', '.join(str(order.index(key)+1) for key in keys)+']'
    source=re.sub(r'\[(@[^\]]+)\]',cite,source)
    assert set(order)==set(refs),(set(refs)-set(order))
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(1);sec.footer_distance=Inches(.42)
    for el in list(doc.styles.element.iter()):
        if el.tag==qn('w:pBdr'):el.getparent().remove(el)
        elif el.tag==qn('w:rFonts'):
            for attr in list(el.attrib):
                if attr.endswith('Theme'):del el.attrib[attr]
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5);normal.font.color.rgb=RGBColor.from_string('111111')
    normal.paragraph_format.line_spacing=1.14;normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.widow_control=True
    for name,size in [('Title',23),('Subtitle',13),('Heading 1',13)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('111111');style.font.bold=name!='Subtitle';style.font.italic=False
        style.paragraph_format.space_before=Pt(12 if name=='Heading 1' else 0);style.paragraph_format.space_after=Pt(7);style.paragraph_format.keep_with_next=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    doc.core_properties.title='Job seeking is a public health issue';doc.core_properties.subject='Empirical opinion editorial: the hiring environment as a potential psychiatric exposure'
    doc.core_properties.author='Jacob E. Thomas';doc.core_properties.keywords='job seeking; labor market; mental distress; public health; opinion'
    doc.core_properties.created=doc.core_properties.modified=datetime.datetime(2026,9,30)
    front=True;endnote=False;md=[]
    for block in source.strip().split('\n\n'):
        if block=='@@BODY@@':front=False;continue
        if block=='@@REFERENCES@@':
            doc.add_page_break();doc.add_paragraph('References','Heading 1');md.append('## References')
            for number,key in enumerate(order,1):
                ref=refs[key];p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.23);p.paragraph_format.first_line_indent=Inches(-.23);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
                run=p.add_run(f'{number}. '+ref['text']+' ');run.font.size=Pt(10.5);hyperlink(p,ref['url'],ref['url'])
                md.append(f'{number}. '+ref['text']+' '+ref['url'])
            continue
        assert not block.startswith('@@'),block
        md.append(block)
        if block.startswith('# '):doc.add_paragraph(block[2:],'Title');continue
        if block.startswith('## '):
            if block=='## About the supporting analysis':endnote=True
            doc.add_paragraph(block[3:],'Subtitle' if front else 'Heading 1');continue
        p=doc.add_paragraph(block)
        if front:
            p.paragraph_format.space_after=Pt(5 if not block.startswith('Article type:') else 12);p.paragraph_format.keep_with_next=True
            if block.startswith('Article type:'):
                for run in p.runs:run.font.size=Pt(10)
        if endnote:
            p.paragraph_format.line_spacing=1.08
            for run in p.runs:run.font.size=Pt(10.5)
    target=ROOT/'manuscript/Job_seeking_as_a_psychiatric_exposure.docx';memory=io.BytesIO();doc.save(memory)
    with zipfile.ZipFile(memory) as original,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as out:
        for name in original.namelist():
            info=zipfile.ZipInfo(name,date_time=(2026,9,30,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;out.writestr(info,original.read(name))
    (ROOT/'manuscript/manuscript.md').write_text('\n\n'.join(md)+'\n',encoding='utf8',newline='\n')
    (ROOT/'audit/citation_order.json').write_text(json.dumps(order,indent=2)+'\n',encoding='utf8',newline='\n')
    counts={'main_argument_words':len(source.split('@@BODY@@')[1].split('## About the supporting analysis')[0].split()),'all_text_including_references_words':len(' '.join(md).split()),'references':len(order),'figures':0,'tables':0}
    (ROOT/'audit/manuscript_counts.json').write_text(json.dumps(counts,indent=2)+'\n',encoding='utf8',newline='\n')
    print('Built empirical opinion editorial:',json.dumps(counts))

if __name__=='__main__':main()
